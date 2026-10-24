import os
import sys

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from build_report.helpers import set_cell_margins, set_cell_borders
from build_report.preliminary import (
    add_cover_page, add_certificate, add_declaration, add_acknowledgement,
    add_abstract, add_table_of_contents, add_list_of_figures, add_list_of_tables,
    add_list_of_abbreviations
)
from build_report.chapter1 import add_chapter_1
from build_report.chapter2 import add_chapter_2
from build_report.chapter3 import add_chapter_3
from build_report.chapter4 import add_chapter_4
from build_report.chapter5 import add_chapter_5
from build_report.chapter6 import add_chapter_6
from build_report.chapter7 import add_chapter_7
from build_report.references_appendix import add_references_and_appendix

def setup_page_geometry(section):
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)

def main():
    print("=== BUILDING DEEPGUARD AI MINI PROJECT REPORT ===")
    doc = docx.Document()

    # SECTION 1: Preliminary Pages (No Header, No Footer)
    s1 = doc.sections[0]
    setup_page_geometry(s1)

    print("1/10 Adding Cover Page...")
    add_cover_page(doc)

    print("2/10 Adding Certificate...")
    add_certificate(doc)

    print("3/10 Adding Declaration...")
    add_declaration(doc)

    print("4/10 Adding Acknowledgement...")
    add_acknowledgement(doc)

    print("5/10 Adding Abstract...")
    add_abstract(doc)

    print("6/10 Adding Table of Contents...")
    add_table_of_contents(doc)

    print("7/10 Adding List of Figures...")
    add_list_of_figures(doc)

    print("8/10 Adding List of Tables...")
    add_list_of_tables(doc)

    print("9/10 Adding List of Abbreviations...")
    add_list_of_abbreviations(doc)

    # SECTION 2: Chapter Pages Onward (Header and Footer Active)
    print("10/10 Configuring Section 2 (Chapter Pages Onward)...")
    s2 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
    setup_page_geometry(s2)

    # Unlink headers and footers from Section 1
    s2.header.is_linked_to_previous = False
    s2.footer.is_linked_to_previous = False

    # Reset page numbering to start at 1 in Section 2
    pgNumType = parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1"/>')
    s2._sectPr.append(pgNumType)

    # HEADER: Right-aligned "Dual-Stream Deepfake Video Detection | Mini Project Report"
    hp = s2.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("Dual-Stream Deepfake Video Detection  |  Mini Project Report")
    hrun.font.name = 'Times New Roman'
    hrun.font.size = Pt(9.5)
    hrun.font.italic = True
    hrun.font.color.rgb = RGBColor(113, 128, 150)

    # FOOTER: Center-aligned page number ONLY (No institutional text)
    p_footer = s2.footer.paragraphs[0]
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_num = p_footer.add_run()
    r_num.font.name = 'Times New Roman'
    r_num.font.size = Pt(10)
    r_num.font.color.rgb = RGBColor(45, 55, 72)
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    r_num._r.append(fldChar1)
    r_num._r.append(instrText)
    r_num._r.append(fldChar2)
    r_num._r.append(fldChar3)

    print("Compiling Chapter 1...")
    add_chapter_1(doc)

    print("Compiling Chapter 2...")
    add_chapter_2(doc)

    print("Compiling Chapter 3...")
    add_chapter_3(doc)

    print("Compiling Chapter 4...")
    add_chapter_4(doc)

    print("Compiling Chapter 5...")
    add_chapter_5(doc)

    print("Compiling Chapter 6...")
    add_chapter_6(doc)

    print("Compiling Chapter 7...")
    add_chapter_7(doc)

    print("Compiling References & Appendix...")
    add_references_and_appendix(doc)

    out_dirs = [r'D:\mini-project\report', r'D:\mini-project\SUBMISSION']
    for d in out_dirs:
        os.makedirs(d, exist_ok=True)
        docx_file = os.path.join(d, 'Mini_Project_Report_Dev_Rana_Kaila.docx')
        doc.save(docx_file)
        print(f"Saved DOCX report to: {docx_file}")

    print("Converting to PDF via Word Automation...")
    import win32com.client
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        primary_docx = os.path.abspath(r'D:\mini-project\report\Mini_Project_Report_Dev_Rana_Kaila.docx')
        primary_pdf = os.path.abspath(r'D:\mini-project\report\Mini_Project_Report_Dev_Rana_Kaila.pdf')
        sub_pdf = os.path.abspath(r'D:\mini-project\SUBMISSION\Mini_Project_Report_Dev_Rana_Kaila.pdf')

        wb = word.Documents.Open(primary_docx, ReadOnly=True)
        wb.SaveAs(primary_pdf, FileFormat=17)
        wb.SaveAs(sub_pdf, FileFormat=17)
        wb.Close()
        print(f"Report PDF successfully generated at: {primary_pdf}")
        print(f"Report PDF copied to submission at: {sub_pdf}")
    finally:
        word.Quit()

if __name__ == '__main__':
    main()
