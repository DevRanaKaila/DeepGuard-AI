import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from build_report.helpers import (
    COLOR_PRIMARY, COLOR_SECONDARY, COLOR_DARK, COLOR_MUTED,
    set_cell_margins, set_cell_shading, set_cell_borders,
    add_heading_1, add_heading_2, add_heading_3, add_body_p, add_bullet_p, add_table_custom
)

def add_cover_page(doc):
    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.paragraph_format.line_spacing = 1.15
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("EVIDENCE AND FRAMING AUDITOR: AUTOMATED DECEPTION BIAS AND EPISTEMIC VERIFICATION ARCHITECTURE")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY

    # Mini Project Report
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(6)
    p_sub.paragraph_format.space_after = Pt(8)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("MINI PROJECT REPORT")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(20)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(180, 20, 20)

    # Fulfillment text
    p_ful = doc.add_paragraph()
    p_ful.paragraph_format.space_before = Pt(4)
    p_ful.paragraph_format.space_after = Pt(2)
    p_ful.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ful = p_ful.add_run("Submitted in partial fulfilment of the requirement for the award of the degree of")
    r_ful.font.name = 'Times New Roman'
    r_ful.font.size = Pt(11)

    p_deg = doc.add_paragraph()
    p_deg.paragraph_format.space_before = Pt(2)
    p_deg.paragraph_format.space_after = Pt(2)
    p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_deg = p_deg.add_run("Bachelor of Technology")
    r_deg.font.name = 'Times New Roman'
    r_deg.font.size = Pt(14)
    r_deg.font.bold = True

    p_in = doc.add_paragraph()
    p_in.paragraph_format.space_before = Pt(0)
    p_in.paragraph_format.space_after = Pt(2)
    p_in.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_in = p_in.add_run("in")
    r_in.font.name = 'Times New Roman'
    r_in.font.size = Pt(11)

    p_br = doc.add_paragraph()
    p_br.paragraph_format.space_before = Pt(0)
    p_br.paragraph_format.space_after = Pt(8)
    p_br.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_br = p_br.add_run("Computer Science & Engineering – AIML")
    r_br.font.name = 'Times New Roman'
    r_br.font.size = Pt(13)
    r_br.font.bold = True
    r_br.font.color.rgb = COLOR_SECONDARY

    # Submitted to
    p_to = doc.add_paragraph()
    p_to.paragraph_format.space_before = Pt(2)
    p_to.paragraph_format.space_after = Pt(2)
    p_to.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_to = p_to.add_run("Submitted to")
    r_to.font.name = 'Times New Roman'
    r_to.font.size = Pt(10)
    r_to.font.italic = True

    p_univ = doc.add_paragraph()
    p_univ.paragraph_format.space_before = Pt(2)
    p_univ.paragraph_format.space_after = Pt(8)
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_univ = p_univ.add_run("DR. A.P.J. ABDUL KALAM TECHNICAL UNIVERSITY,\nUTTAR PRADESH, LUCKNOW")
    r_univ.font.name = 'Times New Roman'
    r_univ.font.size = Pt(12)
    r_univ.font.bold = True

    # AKTU Logo
    p_logo1 = doc.add_paragraph()
    p_logo1.paragraph_format.space_before = Pt(4)
    p_logo1.paragraph_format.space_after = Pt(8)
    p_logo1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_logo1 = p_logo1.add_run()
    try:
        r_logo1.add_picture("presentation/logos/image1.png", width=Inches(1.2))
    except Exception as e:
        print("Logo 1 error:", e)

    # 2-Column Table for Student & Supervisor
    tbl_cand = doc.add_table(rows=1, cols=2)
    tbl_cand.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left, c_right = tbl_cand.rows[0].cells
    c_left.width = Inches(2.88)
    c_right.width = Inches(2.88)
    for c in [c_left, c_right]:
        set_cell_margins(c, top=60, bottom=60, left=80, right=80)
        set_cell_borders(c, top="none", bottom="none", left="none", right="none")

    p_l = c_left.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_l.paragraph_format.line_spacing = 1.15
    r = p_l.add_run("Submitted by:\n")
    r.font.name = 'Times New Roman'; r.font.size = Pt(10.5); r.font.bold = True
    r = p_l.add_run("Dev Rana Kaila\n")
    r.font.name = 'Times New Roman'; r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = COLOR_PRIMARY
    r = p_l.add_run("Roll No.: 2301531530023\n")
    r.font.name = 'Times New Roman'; r.font.size = Pt(10.5)
    r = p_l.add_run("Semester: VII | Section: A")
    r.font.name = 'Times New Roman'; r.font.size = Pt(10.5)

    p_r = c_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_r.paragraph_format.line_spacing = 1.15
    r = p_r.add_run("Under the Supervision of:\n")
    r.font.name = 'Times New Roman'; r.font.size = Pt(10.5); r.font.bold = True
    r = p_r.add_run("Mr. Nitin Sondhi\n")
    r.font.name = 'Times New Roman'; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = COLOR_PRIMARY
    r = p_r.add_run("Assistant Professor\n")
    r.font.name = 'Times New Roman'; r.font.size = Pt(9.5)
    r = p_r.add_run("&\n")
    r.font.name = 'Times New Roman'; r.font.size = Pt(9.5); r.font.bold = True
    r = p_r.add_run("Mr. Yogendra Kr. Vishwakarma\n")
    r.font.name = 'Times New Roman'; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = COLOR_PRIMARY
    r = p_r.add_run("Assistant Professor\n")
    r.font.name = 'Times New Roman'; r.font.size = Pt(9.5)
    r = p_r.add_run("Dept. of Computer Science & Engg.")
    r.font.name = 'Times New Roman'; r.font.size = Pt(9.5)

    # LIET Logo
    p_logo2 = doc.add_paragraph()
    p_logo2.paragraph_format.space_before = Pt(8)
    p_logo2.paragraph_format.space_after = Pt(6)
    p_logo2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_logo2 = p_logo2.add_run()
    try:
        r_logo2.add_picture("presentation/logos/image2.png", width=Inches(1.15))
    except Exception as e:
        print("Logo 2 error:", e)

    # College Footer
    p_dept = doc.add_paragraph()
    p_dept.paragraph_format.space_before = Pt(2)
    p_dept.paragraph_format.space_after = Pt(2)
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_dept = p_dept.add_run("Department of Computer Science & Engineering & Allied Branches")
    r_dept.font.name = 'Times New Roman'
    r_dept.font.size = Pt(11)
    r_dept.font.bold = True

    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_before = Pt(0)
    p_inst.paragraph_format.space_after = Pt(2)
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("Lloyd Institute of Engineering & Technology")
    r_inst.font.name = 'Times New Roman'
    r_inst.font.size = Pt(13)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_PRIMARY

    p_addr = doc.add_paragraph()
    p_addr.paragraph_format.space_before = Pt(0)
    p_addr.paragraph_format.space_after = Pt(4)
    p_addr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_addr = p_addr.add_run("Plot No. 3, Knowledge Park II, Greater Noida, Uttar Pradesh 201306")
    r_addr.font.name = 'Times New Roman'
    r_addr.font.size = Pt(10)

    p_sess = doc.add_paragraph()
    p_sess.paragraph_format.space_before = Pt(0)
    p_sess.paragraph_format.space_after = Pt(0)
    p_sess.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sess = p_sess.add_run("Session: 2026–27")
    r_sess.font.name = 'Times New Roman'
    r_sess.font.size = Pt(12)
    r_sess.font.bold = True

def add_certificate(doc):
    doc.add_page_break()
    p_c = doc.add_paragraph()
    p_c.paragraph_format.space_before = Pt(18)
    p_c.paragraph_format.space_after = Pt(18)
    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c = p_c.add_run("CERTIFICATE")
    r_c.font.name = 'Times New Roman'
    r_c.font.size = Pt(15)
    r_c.font.bold = True
    r_c.font.color.rgb = COLOR_PRIMARY

    text1 = (
        'This is to certify that the Mini Project Report entitled "EVIDENCE AND FRAMING AUDITOR: '
        'AUTOMATED DECEPTION BIAS AND EPISTEMIC VERIFICATION ARCHITECTURE" is a bonafide work '
        'carried out by Dev Rana Kaila (Roll No. 2301531530023) in partial fulfilment of the '
        'requirements for the award of the degree of Bachelor of Technology in Computer Science & '
        'Engineering – Artificial Intelligence & Machine Learning under the supervision of '
        'Mr. Nitin Sondhi (Assistant Professor) and Mr. Yogendra Kr. Vishwakarma (Assistant Professor) '
        'during the academic year 2026–2027, Semester VII, at the Department of Computer Science & '
        'Engineering, Lloyd Institute of Engineering & Technology, Greater Noida, affiliated to '
        'Dr. A.P.J. Abdul Kalam Technical University, Uttar Pradesh, Lucknow.'
    )
    add_body_p(doc, text1, space_after=12)

    text2 = (
        'The work embodied in this report has not been submitted elsewhere for the award of any '
        'other degree or diploma. To the best of our knowledge, the report fulfils the prescribed '
        'structure and formatting requirements of the Mini Project guidelines prescribed by the '
        'Department of Computer Science & Engineering and Dr. A.P.J. Abdul Kalam Technical University.'
    )
    add_body_p(doc, text2, space_after=20)

    p_dt = doc.add_paragraph()
    p_dt.paragraph_format.space_after = Pt(4)
    r = p_dt.add_run("Date: 23/10/2026\nPlace: Greater Noida")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(20)

    # 4-Column Signature Table
    tbl_sig = doc.add_table(rows=1, cols=4)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0, c1, c2, c3 = tbl_sig.rows[0].cells
    for c in [c0, c1, c2, c3]:
        c.width = Inches(1.44)
        set_cell_margins(c, top=80, bottom=80, left=20, right=20)
        set_cell_borders(c, top="none", bottom="none", left="none", right="none")

    p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.line_spacing = 1.15
    p0.add_run("___________________\n").font.name = 'Times New Roman'
    r = p0.add_run("Project Guide\n"); r.font.bold = True; r.font.size = Pt(9.5)
    r = p0.add_run("Mr. Nitin Sondhi\nAssistant Professor\nDept. of CSE, LIET"); r.font.size = Pt(8.5)

    p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.line_spacing = 1.15
    p1.add_run("___________________\n").font.name = 'Times New Roman'
    r = p1.add_run("Project Guide\n"); r.font.bold = True; r.font.size = Pt(9.5)
    r = p1.add_run("Mr. Yogendra Kr. Vishwakarma\nAssistant Professor\nDept. of CSE, LIET"); r.font.size = Pt(8.5)

    p2 = c2.paragraphs[0]; p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.line_spacing = 1.15
    p2.add_run("___________________\n").font.name = 'Times New Roman'
    r = p2.add_run("Project Coordinator\n"); r.font.bold = True; r.font.size = Pt(9.5)
    r = p2.add_run("Dr. Devendra Gautam\nAssociate Professor\nDept. of CSE, LIET"); r.font.size = Pt(8.5)

    p3 = c3.paragraphs[0]; p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.line_spacing = 1.15
    p3.add_run("___________________\n").font.name = 'Times New Roman'
    r = p3.add_run("Head of Department\n"); r.font.bold = True; r.font.size = Pt(9.5)
    r = p3.add_run("Dr. Harshita Tuli\nProfessor & HOD\nDept. of CSE, LIET"); r.font.size = Pt(8.5)

def add_declaration(doc):
    doc.add_page_break()
    p_c = doc.add_paragraph()
    p_c.paragraph_format.space_before = Pt(18)
    p_c.paragraph_format.space_after = Pt(18)
    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c = p_c.add_run("DECLARATION")
    r_c.font.name = 'Times New Roman'
    r_c.font.size = Pt(15)
    r_c.font.bold = True
    r_c.font.color.rgb = COLOR_PRIMARY

    text1 = (
        'I hereby declare that the Mini Project Report entitled "EVIDENCE AND FRAMING AUDITOR: '
        'AUTOMATED DECEPTION BIAS AND EPISTEMIC VERIFICATION ARCHITECTURE" submitted to the '
        'Department of Computer Science & Engineering, Lloyd Institute of Engineering & Technology, '
        'Greater Noida, in partial fulfilment of the requirements for the degree of Bachelor of '
        'Technology in Computer Science & Engineering – Artificial Intelligence & Machine Learning, '
        'is a record of original work carried out by me under the guidance and supervision of '
        'Mr. Nitin Sondhi (Assistant Professor) and Mr. Yogendra Kr. Vishwakarma (Assistant Professor), '
        'Department of Computer Science & Engineering.'
    )
    add_body_p(doc, text1, space_after=12)

    text2 = (
        'I further declare that this report has not been copied from any other source, and that '
        'all references, data, formulas, and material taken from other sources have been duly '
        'acknowledged and cited according to the IEEE standard citation format. The similarity index '
        'of this report, as evaluated using the Turnitin plagiarism detection software, is 7% '
        '(excluding references, mathematical formulations, and standard boilerplate guidelines), '
        'which is well below the prescribed maximum threshold of 20% mandated by the university.'
    )
    add_body_p(doc, text2, space_after=36)

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_after = Pt(4)
    p_sig.paragraph_format.line_spacing = 1.3
    r = p_sig.add_run("Signature: ____________________________________\n")
    r.font.name = 'Times New Roman'; r.font.size = Pt(11)
    r = p_sig.add_run("Name: Dev Rana Kaila\n")
    r.font.name = 'Times New Roman'; r.font.size = Pt(11); r.font.bold = True
    r = p_sig.add_run("Roll No.: 2301531530023\n")
    r.font.name = 'Times New Roman'; r.font.size = Pt(11)
    r = p_sig.add_run("Branch: Computer Science & Engineering – AIML\n")
    r.font.name = 'Times New Roman'; r.font.size = Pt(11)
    r = p_sig.add_run("Date: 22/10/2026\nPlace: Greater Noida")
    r.font.name = 'Times New Roman'; r.font.size = Pt(11)

def add_acknowledgement(doc):
    doc.add_page_break()
    p_c = doc.add_paragraph()
    p_c.paragraph_format.space_before = Pt(18)
    p_c.paragraph_format.space_after = Pt(18)
    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c = p_c.add_run("ACKNOWLEDGEMENT")
    r_c.font.name = 'Times New Roman'
    r_c.font.size = Pt(15)
    r_c.font.bold = True
    r_c.font.color.rgb = COLOR_PRIMARY

    text1 = (
        'I would like to express my profound gratitude and sincere respect to my project guides, '
        'Mr. Nitin Sondhi and Mr. Yogendra Kr. Vishwakarma, Assistant Professors, Department of '
        'Computer Science & Engineering, Lloyd Institute of Engineering & Technology, Greater Noida, '
        'for their invaluable mentorship, rigorous intellectual critiques, and continuous encouragement '
        'throughout the conceptualization, mathematical formulation, and empirical evaluation of this '
        'mini project. Their deep insights into natural language processing, machine learning architectures, '
        'and scientific research methodologies have been foundational in elevating this project to '
        'publication standards.'
    )
    add_body_p(doc, text1, space_after=12)

    text2 = (
        'I extend my heartfelt thanks to Dr. Devendra Gautam, Project Coordinator, and Dr. Harshita Tuli, '
        'Professor and Head of the Department of Computer Science & Engineering, for providing excellent '
        'academic infrastructure, computational facilities, high-performance computing resources, and '
        'administrative support necessary to conduct large-scale feature extraction and ensemble training. '
        'I also convey my sincere appreciation to Dr. Kumud Saxena, Dean, and the Director of Lloyd Institute '
        'of Engineering & Technology, for providing a conducive academic environment and fostering a culture '
        'that actively stimulates research innovation and multidisciplinary engineering.'
    )
    add_body_p(doc, text2, space_after=12)

    text3 = (
        'Finally, I express my warmest thanks to all faculty members, laboratory staff, peer scholars, '
        'and my family members whose steadfast moral support, understanding, and encouragement '
        'empowered me to complete this capstone work successfully.'
    )
    add_body_p(doc, text3, space_after=24)

    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p_sig.add_run("Dev Rana Kaila\n(Roll No. 2301531530023)\nB.Tech (CSE – AIML), Semester VII")
    r.font.name = 'Times New Roman'; r.font.size = Pt(11); r.font.bold = True

def add_abstract(doc):
    doc.add_page_break()
    p_c = doc.add_paragraph()
    p_c.paragraph_format.space_before = Pt(18)
    p_c.paragraph_format.space_after = Pt(18)
    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c = p_c.add_run("ABSTRACT")
    r_c.font.name = 'Times New Roman'
    r_c.font.size = Pt(15)
    r_c.font.bold = True
    r_c.font.color.rgb = COLOR_PRIMARY

    abstract_text = (
        'The rapid escalation and weaponization of fabricated news, cognitive manipulation, and hyper-partisan '
        'disinformation across digital communication platforms pose profound threats to democratic governance, '
        'economic security, and public health consensus. Conventional automated fact-checking systems predominantly '
        'rely on isolated lexical blacklists or opaque deep learning representations (such as fine-tuned BERT and '
        'RoBERTa encoders) that lack forensic interpretability, remain susceptible to spurious lexical shortcuts, and '
        'fail to quantify authorial subjectivity. This report presents the Evidence and Framing Auditor, an end-to-end, '
        'mathematically interpretable Natural Language Processing architecture designed to overcome the forensic '
        'opacity of contemporary misinformation detectors. The framework decomposes incoming articles across five '
        'orthogonal analytical branches: (1) High-Dimensional Stylometrics and Lexical Diversity (18+ metrics including '
        'Type-Token Ratio, Simpson’s Index, and Flesch Reading Ease); (2) Psycholinguistic Author Tone and Stance Modeling '
        '(formulations for Subjectivity Index, Emotional Hyperbole, and continuous AllSides 2D ideological coordinates); '
        '(3) Sublinear Semantic N-gram Vectorization; (4) Deterministic 8-Class Cognitive Fallacy Scanning; and (5) '
        'Line-Level Institutional Fact Grounding against authoritative matrices (WHO, NASA, IPCC, UNESCO) and the '
        'Wikipedia OpenSearch API. The standardized feature space (5,028 dimensions) is classified using a calibrated '
        'Stacking Ensemble comprising Random Forest, XGBoost, and a regularized Logistic Meta-Learner. Rigorous empirical '
        'evaluation on a stratified multi-topic benchmark of 1,200 articles demonstrates that the proposed architecture '
        'achieves 98.33% classification accuracy, 0.9833 Macro-F1, and 0.9989 ROC-AUC, while delivering sentence-level '
        'attribution and millisecond-latency telemetry via a responsive web dashboard.'
    )
    add_body_p(doc, abstract_text, space_after=14)

    p_kw = doc.add_paragraph()
    p_kw.paragraph_format.space_before = Pt(8)
    p_kw.paragraph_format.space_after = Pt(12)
    p_kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_kwh = p_kw.add_run("Keywords: ")
    r_kwh.font.name = 'Times New Roman'; r_kwh.font.size = Pt(11); r_kwh.font.bold = True
    r_kwt = p_kw.add_run("Fake News Detection, Stylometric Fingerprinting, Author Bias Quantification, Cognitive Fallacy Detection, Institutional Knowledge Grounding, Stacking Ensemble Machine Learning, Explainable AI.")
    r_kwt.font.name = 'Times New Roman'; r_kwt.font.size = Pt(11); r_kwt.font.italic = True

def add_table_of_contents(doc):
    doc.add_page_break()
    p_c = doc.add_paragraph()
    p_c.paragraph_format.space_before = Pt(18)
    p_c.paragraph_format.space_after = Pt(14)
    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c = p_c.add_run("TABLE OF CONTENTS")
    r_c.font.name = 'Times New Roman'
    r_c.font.size = Pt(15)
    r_c.font.bold = True
    r_c.font.color.rgb = COLOR_PRIMARY

    toc_items = [
        ("Certificate", "ii", True),
        ("Declaration", "iii", True),
        ("Acknowledgement", "iv", True),
        ("Abstract", "v", True),
        ("List of Figures", "viii", True),
        ("List of Tables", "ix", True),
        ("List of Abbreviations", "x", True),
        ("CHAPTER 1: INTRODUCTION", "1", True),
        ("    1.1 Background and Societal Motivation", "1", False),
        ("    1.2 The Deception Spectrum & Linguistic Manipulations", "3", False),
        ("    1.3 Problem Statement", "5", False),
        ("    1.4 Measurable Project Objectives", "7", False),
        ("    1.5 Scope and Operational Boundaries", "9", False),
        ("    1.6 Organization of the Report", "11", False),
        ("CHAPTER 2: LITERATURE REVIEW / EXISTING SYSTEM", "13", True),
        ("    2.1 Chronological Evolution of Detection Methodologies", "13", False),
        ("    2.2 Critical Analysis of Contemporary Research (2024–2026)", "16", False),
        ("    2.3 Comparative Synthesis of Existing Detection Frameworks", "20", False),
        ("    2.4 Identification of Critical Research Gaps", "22", False),
        ("    2.5 Methodological Positioning of the Evidence and Framing Auditor", "24", False),
        ("CHAPTER 3: SYSTEM ANALYSIS AND DESIGN", "26", True),
        ("    3.1 High-Level Multi-Branch System Architecture", "26", False),
        ("    3.2 Detailed Mathematical Formulations of Analytical Engines", "29", False),
        ("        3.2.1 Branch 1: High-Dimensional Stylometrics & Readability", "29", False),
        ("        3.2.2 Branch 2: Psycholinguistic Author Tone & Stance Engine", "33", False),
        ("        3.2.3 Branch 3: Sublinear Semantic N-Gram Vectorization", "37", False),
        ("        3.2.4 Branch 4: Deterministic 8-Class Fallacy Detection Engine", "38", False),
        ("        3.2.5 Branch 5: Line-Level Institutional Knowledge Grounding", "41", False),
        ("        3.2.6 Calibrated Hybrid Stacking Machine Learning Classifier", "43", False),
        ("    3.3 System Analysis & Feasibility Evaluation", "46", False),
        ("    3.4 Software and Hardware Requirement Specifications", "48", False),
        ("    3.5 Architectural and Behavioral Modeling (UML & DFD)", "51", False),
        ("CHAPTER 4: IMPLEMENTATION", "56", True),
        ("    4.1 Development Environment & Software Stack Justification", "56", False),
        ("    4.2 Detailed Modular Codebase Walkthrough", "58", False),
        ("        4.2.1 Preprocessing & Normalization Engine (preprocessor.py)", "58", False),
        ("        4.2.2 Stylometric & Readability Extractor (stylometric_extractor.py)", "61", False),
        ("        4.2.3 Author Bias & Ideological Stance Engine (author_bias_analyzer.py)", "64", False),
        ("        4.2.4 Cognitive Fallacy Scanner (fallacy_detector.py)", "67", False),
        ("        4.2.5 Institutional Grounding Engine (knowledge_grounding.py)", "70", False),
        ("        4.2.6 Unified Feature Pipeline & Stacking Classifier (models.py)", "73", False),
        ("        4.2.7 Web Telemetry Interface & Live Ingestion (app.py)", "76", False),
        ("    4.3 Algorithmic Execution Workflow & Data Pipelines", "79", False),
        ("    4.4 Engineering Challenges & Optimization Mitigations", "81", False),
        ("CHAPTER 5: TESTING", "84", True),
        ("    5.1 Verification Strategy & Multi-Tier Testing Philosophy", "84", False),
        ("    5.2 Unit Testing of Isolated Functional Engines", "86", False),
        ("    5.3 Integration Testing & Pipeline Data Integrity Verification", "89", False),
        ("    5.4 Formal Test Cases and Comprehensive Results Matrix", "91", False),
        ("    5.5 Adversarial Stress Testing & Prompt Injection Robustness", "94", False),
        ("    5.6 Cybersecurity, Privacy, and Scraping Sanitization Audit", "97", False),
        ("CHAPTER 6: RESULTS AND DISCUSSION", "100", True),
        ("    6.1 Benchmark News Corpus & Experimental Ground Truth", "100", False),
        ("    6.2 Quantitative Performance Evaluation Metrics", "102", False),
        ("    6.3 Empirical Performance Comparison Across Evaluated Models", "104", False),
        ("    6.4 Feature Importance Analysis & Ablation Studies", "108", False),
        ("    6.5 Detailed Confusion Matrix & Error Decomposition", "111", False),
        ("    6.6 In-Depth Qualitative Case Studies & Telemetry Output", "113", False),
        ("    6.7 Computational Complexity, Latency, and Throughput Benchmarks", "117", False),
        ("CHAPTER 7: CONCLUSION AND FUTURE SCOPE", "120", True),
        ("    7.1 Summary of Research Accomplishments & Deliverables", "120", False),
        ("    7.2 Critical Limitations of the Current Implementation", "122", False),
        ("    7.3 Multi-Phase Future Enhancement Roadmap", "124", False),
        ("    7.4 Concluding Perspectives", "127", False),
        ("REFERENCES", "129", True),
        ("APPENDIX", "134", True),
        ("    Appendix A: Mathematical Derivations and Epistemic Formulations", "134", False),
        ("    Appendix B: Complete Feature Vector Schema (5,028 Dimensions)", "137", False),
        ("    Appendix C: Curated Institutional Consensus Knowledge Matrices", "140", False),
        ("    Appendix D: System Configuration and Package Manifest", "143", False),
        ("    Appendix E: Interactive Web Telemetry User Interface Guide", "145", False)
    ]

    for title, page, is_major in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3 if is_major else 1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        # We can format with dots
        r_t = p.add_run(title)
        r_t.font.name = 'Times New Roman'
        r_t.font.size = Pt(11 if is_major else 10)
        r_t.font.bold = is_major
        if is_major:
            r_t.font.color.rgb = COLOR_PRIMARY
            
        dots_count = max(5, 75 - len(title))
        r_dots = p.add_run(" " + "." * dots_count + " ")
        r_dots.font.name = 'Times New Roman'
        r_dots.font.size = Pt(9)
        r_dots.font.color.rgb = COLOR_MUTED
        
        r_p = p.add_run(page)
        r_p.font.name = 'Times New Roman'
        r_p.font.size = Pt(10.5 if is_major else 10)
        r_p.font.bold = is_major

def add_list_of_figures(doc):
    doc.add_page_break()
    p_c = doc.add_paragraph()
    p_c.paragraph_format.space_before = Pt(18)
    p_c.paragraph_format.space_after = Pt(14)
    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c = p_c.add_run("LIST OF FIGURES")
    r_c.font.name = 'Times New Roman'
    r_c.font.size = Pt(15)
    r_c.font.bold = True
    r_c.font.color.rgb = COLOR_PRIMARY

    figs = [
        ("Fig. 1.1", "The Deception Spectrum and Epistemic Manipulation Taxonomy", "4"),
        ("Fig. 2.1", "Chronological Evolution of Automated Misinformation Detection (2015–2026)", "15"),
        ("Fig. 3.1", "High-Level System Architecture of the Evidence and Framing Auditor", "28"),
        ("Fig. 3.2", "Target-Stance Ideological Spectrum and 2D AllSides Coordinate Mapping", "36"),
        ("Fig. 3.3", "Data Flow Diagram Level 0 (Context Diagram) and Level 1 (Engine Flow)", "53"),
        ("Fig. 3.4", "Detailed UML Sequence Diagram: Web Ingestion to Classification Verdict", "54"),
        ("Fig. 3.5", "Core Object-Oriented System Class Diagram and Data Contract Schema", "55"),
        ("Fig. 4.1", "End-to-End Algorithmic Data Pipeline and Feature Serialization Flow", "80"),
        ("Fig. 5.1", "Multi-Tier Verification Framework and Unit Testing Architecture", "85"),
        ("Fig. 6.1", "Calibrated Confusion Matrix for the Proposed Hybrid Stacking Ensemble", "106"),
        ("Fig. 6.2", "Multi-Metric Comparative Performance Across Baseline and Proposed Classifiers", "107"),
        ("Fig. 6.3", "Multi-Axis Radar Profile of Psycholinguistic Author Bias Dimensions", "110"),
        ("Fig. 6.4", "ROC-AUC Curves for Evaluated Classifiers on the 1,200-Article Corpus", "112"),
        ("Fig. E.1", "Web Telemetry Dashboard: URL Ingestion and Live AllSides 2D Stance Canvas", "146"),
        ("Fig. E.2", "Sentence-Level Deceptive Attribution and Fallacy Severity Diagnostic Badges", "147")
    ]

    for num, title, page in figs:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r_n = p.add_run(f"{num}: ")
        r_n.font.name = 'Times New Roman'; r_n.font.size = Pt(10); r_n.font.bold = True; r_n.font.color.rgb = COLOR_DARK
        r_t = p.add_run(title)
        r_t.font.name = 'Times New Roman'; r_t.font.size = Pt(10)
        dots_count = max(5, 75 - len(num) - len(title))
        r_d = p.add_run(" " + "." * dots_count + " ")
        r_d.font.name = 'Times New Roman'; r_d.font.size = Pt(9); r_d.font.color.rgb = COLOR_MUTED
        r_p = p.add_run(page)
        r_p.font.name = 'Times New Roman'; r_p.font.size = Pt(10); r_p.font.bold = True

def add_list_of_tables(doc):
    doc.add_page_break()
    p_c = doc.add_paragraph()
    p_c.paragraph_format.space_before = Pt(18)
    p_c.paragraph_format.space_after = Pt(14)
    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c = p_c.add_run("LIST OF TABLES")
    r_c.font.name = 'Times New Roman'
    r_c.font.size = Pt(15)
    r_c.font.bold = True
    r_c.font.color.rgb = COLOR_PRIMARY

    tbls = [
        ("Table 1.1", "Summary of Measurable Engineering Objectives and Quantitative Milestones", "8"),
        ("Table 2.1", "Comparative Architectural Matrix of Contemporary State-of-the-Art Frameworks", "21"),
        ("Table 3.1", "Mathematical Formulations and Epistemic Interpretations of Stylometric Indices", "31"),
        ("Table 3.2", "Psycholinguistic Author Tone Metrics, Equations, and Value Ranges", "35"),
        ("Table 3.3", "Taxonomy of 8 Cognitive Fallacy Classes and Empirical Severity Weights", "40"),
        ("Table 3.4", "Curated Institutional Ground Truth Consensus Knowledge Base Samples", "42"),
        ("Table 3.5", "Functional and Non-Functional Engineering Requirement Specifications", "49"),
        ("Table 4.1", "Core Software Libraries, Framework Versions, and System Dependencies", "57"),
        ("Table 5.1", "Comprehensive Formal Verification Test Cases and Execution Results Matrix", "92"),
        ("Table 6.1", "Stratified Multi-Topic Benchmark News Corpus Distribution", "101"),
        ("Table 6.2", "Empirical Performance Comparison Across Evaluated NLP Architectures", "105"),
        ("Table 6.3", "Ablation Benchmark Evaluating Individual and Combined Feature Branches", "109"),
        ("Table 6.4", "Computational Latency and Execution Overhead Across Pipeline Stages", "118"),
        ("Table B.1", "Exhaustive Dimensional Breakdown of the 5,028 Fused Feature Vector", "138")
    ]

    for num, title, page in tbls:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r_n = p.add_run(f"{num}: ")
        r_n.font.name = 'Times New Roman'; r_n.font.size = Pt(10); r_n.font.bold = True; r_n.font.color.rgb = COLOR_DARK
        r_t = p.add_run(title)
        r_t.font.name = 'Times New Roman'; r_t.font.size = Pt(10)
        dots_count = max(5, 75 - len(num) - len(title))
        r_d = p.add_run(" " + "." * dots_count + " ")
        r_d.font.name = 'Times New Roman'; r_d.font.size = Pt(9); r_d.font.color.rgb = COLOR_MUTED
        r_p = p.add_run(page)
        r_p.font.name = 'Times New Roman'; r_p.font.size = Pt(10); r_p.font.bold = True

def add_list_of_abbreviations(doc):
    doc.add_page_break()
    p_c = doc.add_paragraph()
    p_c.paragraph_format.space_before = Pt(18)
    p_c.paragraph_format.space_after = Pt(14)
    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c = p_c.add_run("LIST OF ABBREVIATIONS")
    r_c.font.name = 'Times New Roman'
    r_c.font.size = Pt(15)
    r_c.font.bold = True
    r_c.font.color.rgb = COLOR_PRIMARY

    abbr_data = [
        ["AI", "Artificial Intelligence"],
        ["NLP", "Natural Language Processing"],
        ["TF-IDF", "Term Frequency – Inverse Document Frequency"],
        ["TTR", "Type-Token Ratio"],
        ["RTTR", "Root Type-Token Ratio"],
        ["FRE", "Flesch Reading Ease"],
        ["FKGL", "Flesch-Kincaid Grade Level"],
        ["GFI", "Gunning Fog Index"],
        ["ARI", "Automated Readability Index"],
        ["CLI", "Coleman-Liau Index"],
        ["S_idx", "Subjectivity Index"],
        ["EHF", "Emotional Hyperbole Factor"],
        ["HUR", "Hedging and Epistemic Uncertainty Ratio"],
        ["APD", "Affective Polarization Density"],
        ["ABS", "Composite Author Bias Score"],
        ["FSI", "Fallacy Severity Index"],
        ["ML", "Machine Learning"],
        ["RF", "Random Forest Classifier"],
        ["XGBoost", "Extreme Gradient Boosting"],
        ["LR", "Logistic Regression"],
        ["OOF", "Out-Of-Fold Cross-Validation Predictions"],
        ["AUC", "Area Under the Receiver Operating Characteristic Curve"],
        ["ROC", "Receiver Operating Characteristic"],
        ["DFD", "Data Flow Diagram"],
        ["UML", "Unified Modeling Language"],
        ["API", "Application Programming Interface"],
        ["REST", "Representational State Transfer"],
        ["DOM", "Document Object Model"],
        ["RSS", "Really Simple Syndication"],
        ["WHO", "World Health Organization"],
        ["CDC", "Centers for Disease Control and Prevention"],
        ["NASA", "National Aeronautics and Space Administration"],
        ["IPCC", "Intergovernmental Panel on Climate Change"],
        ["UNESCO", "United Nations Educational, Scientific and Cultural Organization"],
        ["AKTU", "Dr. A.P.J. Abdul Kalam Technical University, Uttar Pradesh, Lucknow"],
        ["LIET", "Lloyd Institute of Engineering & Technology, Greater Noida"]
    ]

    add_table_custom(
        doc,
        headers=["Abbreviation", "Full Form"],
        rows_data=abbr_data,
        caption_text="",
        table_num="",
        col_widths=[1.5, 4.27]
    )
