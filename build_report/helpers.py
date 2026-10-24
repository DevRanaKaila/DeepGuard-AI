import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

COLOR_PRIMARY = RGBColor(26, 54, 93)     # #1A365D Deep Navy
COLOR_SECONDARY = RGBColor(43, 108, 176) # #2B6CB0 Blue
COLOR_DARK = RGBColor(45, 55, 72)        # #2D3748 Charcoal
COLOR_MUTED = RGBColor(113, 128, 150)    # #718096 Slate

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding in dxa (1 pt = 20 dxa)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="F7FAFC"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", 
                     color="CCCCCC", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    def b_str(side, val):
        if val == "none":
            return f'<w:{side} w:val="none"/>'
        return f'<w:{side} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'{b_str("top", top)}'
        f'{b_str("left", left)}'
        f'{b_str("bottom", bottom)}'
        f'{b_str("right", right)}'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = COLOR_DARK
    return p

def add_body_p(doc, text, bold_prefix="", space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_bullet_p(doc, text, bold_prefix="", level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    return p

def add_equation_p(doc, eq_text, eq_num=""):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_eq = p.add_run(f"    {eq_text}")
    r_eq.font.name = 'Cambria Math'
    r_eq.font.size = Pt(11)
    r_eq.font.italic = True
    if eq_num:
        r_num = p.add_run(f"          ({eq_num})")
        r_num.font.name = 'Times New Roman'
        r_num.font.size = Pt(11)
        r_num.font.bold = True
    return p

def add_figure(doc, img_path, caption_text, fig_num="Fig. 1.1", width_in=5.2):
    p_img = doc.add_paragraph()
    p_img.paragraph_format.space_before = Pt(12)
    p_img.paragraph_format.space_after = Pt(4)
    p_img.paragraph_format.keep_with_next = True
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_in))
    except Exception as e:
        print(f"Error adding picture {img_path}: {e}")
        p_img.add_run(f"[FIGURE NOT FOUND: {img_path}]")
    
    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.space_after = Pt(12)
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_num = p_cap.add_run(f"{fig_num}: ")
    r_num.font.name = 'Times New Roman'
    r_num.font.size = Pt(10)
    r_num.font.bold = True
    r_num.font.color.rgb = COLOR_DARK
    r_txt = p_cap.add_run(caption_text)
    r_txt.font.name = 'Times New Roman'
    r_txt.font.size = Pt(10)
    r_txt.font.italic = True
    return p_cap

def add_table_custom(doc, headers, rows_data, caption_text="", table_num="Table 1.1", col_widths=None):
    if caption_text:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(12)
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.paragraph_format.keep_with_next = True
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_num = p_cap.add_run(f"{table_num}: ")
        r_num.font.name = 'Times New Roman'
        r_num.font.size = Pt(10)
        r_num.font.bold = True
        r_num.font.color.rgb = COLOR_DARK
        r_txt = p_cap.add_run(caption_text)
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(10)
        r_txt.font.italic = True

    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        cell = hdr_cells[i]
        cell.text = title
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        set_cell_shading(cell, "1A365D") # Deep Navy
        set_cell_borders(cell, top="single", bottom="single", left="none", right="none", color="1A365D", sz="8")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for r_idx, row_values in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F7FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = row_cells[c_idx]
            cell.text = str(val)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            set_cell_shading(cell, bg_color)
            set_cell_borders(cell, top="single", bottom="single", left="none", right="none", color="E2E8F0", sz="4")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 or len(str(val)) > 25 else WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(45, 55, 72)

    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_after = Pt(8)
    return table

def add_code_block(doc, code_str, title=""):
    if title:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(8)
        p_t.paragraph_format.space_after = Pt(2)
        p_t.paragraph_format.keep_with_next = True
        r = p_t.add_run(f"Code Listing: {title}")
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = COLOR_DARK

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(5.77)
    set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
    set_cell_shading(cell, "F8F9FA")
    set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="CBD5E0", sz="6")
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(code_str.strip())
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(30, 41, 59)

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_after = Pt(8)
