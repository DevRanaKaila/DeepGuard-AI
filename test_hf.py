import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()
s1 = doc.sections[0]
s1.page_width = Inches(8.27)
s1.page_height = Inches(11.69)
s1.left_margin = Inches(1.5)
s1.right_margin = Inches(1.0)
s1.top_margin = Inches(1.0)
s1.bottom_margin = Inches(1.0)

# Section 1 content (Preliminary)
p1 = doc.add_paragraph('Cover / Preliminary Page')

# Add Section 2 (Chapters)
s2 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
s2.page_width = Inches(8.27)
s2.page_height = Inches(11.69)
s2.left_margin = Inches(1.5)
s2.right_margin = Inches(1.0)
s2.top_margin = Inches(1.0)
s2.bottom_margin = Inches(1.0)

s2.header.is_linked_to_previous = False
s2.footer.is_linked_to_previous = False

# Set page start to 1 in s2
pgNumType = parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1"/>')
s2._sectPr.append(pgNumType)

# Header in s2
hp = s2.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hrun = hp.add_run('Mini Project Report')
hrun.font.name = 'Times New Roman'
hrun.font.size = Pt(10)
hrun.font.italic = True

# Footer in s2: 3-column table
tbl = s2.footer.add_table(1, 3, Inches(5.77))
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for r in tbl.rows:
    for cell in r.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
        tcPr.append(tcBorders)

c0, c1, c2 = tbl.rows[0].cells
c0.width = Inches(1.0)
c1.width = Inches(3.77)
c2.width = Inches(1.0)

p_dept = c1.paragraphs[0]
p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_dept = p_dept.add_run('Department of Computer Science & Engineering')
r_dept.font.name = 'Times New Roman'
r_dept.font.size = Pt(9)

p_num = c2.paragraphs[0]
p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r_num = p_num.add_run()
r_num.font.name = 'Times New Roman'
r_num.font.size = Pt(9)
fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
r_num._r.append(fldChar1)
r_num._r.append(instrText)
r_num._r.append(fldChar2)
r_num._r.append(fldChar3)

doc.add_paragraph('Chapter 1 Page 1')
doc.save('test_header_footer.docx')
print('Successfully saved test_header_footer.docx')
