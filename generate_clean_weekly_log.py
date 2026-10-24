import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def generate_weekly_log():
    doc = docx.Document()
    
    # Page setup: A4 Landscape
    s = doc.sections[0]
    s.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
    s.page_width = Inches(11.693)
    s.page_height = Inches(8.268)
    s.left_margin = Inches(0.7)
    s.right_margin = Inches(0.6)
    s.top_margin = Inches(0.5)
    s.bottom_margin = Inches(0.5)
    
    # Title paragraphs
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run('Lloyd Institute of Engineering & Technology')
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(15)
    r1.font.bold = True
    
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run('Department of Computer Science & Engineering (AIML)  |  Session 2026-27')
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(11)
    
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(2)
    p3.paragraph_format.space_after = Pt(8)
    r3 = p3.add_run('MINI PROJECT - WEEKLY PROGRESS LOG')
    r3.font.name = 'Times New Roman'
    r3.font.size = Pt(13)
    r3.font.bold = True
    
    # Table 0: Metadata Table
    t0 = doc.add_table(5, 4)
    t0.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        [('Student Name', True), ('Dev Rana Kaila', False), ('Roll Number', True), ('2301531530023', False)],
        [('Branch / Semester', True), ('B.Tech CSE (AIML) - Sem VII', False), ('Section', True), ('A', False)],
        [('Project Guide', True), ('Dr. Devendra Gautam', False), ('Designation', True), ('Associate Professor, Dept. of CSE', False)],
        [('Project Title', True), ('Dual-Stream Spatial and Frequency-Domain Deepfake Video Detection', False), ('', False), ('', False)],
        [('GitHub', True), ('github.com/DevRanaKaila/DeepGuard-AI', False), ('College', True), ('LIET, Greater Noida', False)]
    ]
    
    t0_widths = [Inches(1.8), Inches(3.6), Inches(1.8), Inches(3.1)]
    for r_idx, row in enumerate(t0.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.width = t0_widths[c_idx]
            text, bold = meta_data[r_idx][c_idx]
            cell.text = text
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.space_after = Pt(2)
            for run in cp.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(9.5)
                run.font.bold = bold
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="80" w:type="dxa"/><w:bottom w:w="80" w:type="dxa"/><w:left w:w="100" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tcMar>')
            tcPr.append(tcMar)
            borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="888888\"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="888888\"/><w:left w:val="single" w:sz="4" w:space="0" w:color="888888\"/><w:right w:val="single" w:sz="4" w:space="0" w:color="888888\"/></w:tcBorders>')
            tcPr.append(borders)
            if bold:
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F4F8"/>')
                tcPr.append(shd)
                
    # Section Heading
    p_act = doc.add_paragraph()
    p_act.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_act.paragraph_format.space_before = Pt(8)
    p_act.paragraph_format.space_after = Pt(4)
    r_act = p_act.add_run('Weekly Activity Log (Guide Signatures Required)')
    r_act.font.name = 'Times New Roman'
    r_act.font.size = Pt(11)
    r_act.font.bold = True
    
    # Table 1: Activity Log Table (Week 1 removed; replaced by Date Range)
    t1_widths = [Inches(1.85), Inches(1.5), Inches(4.95), Inches(0.9), Inches(1.1)]
    
    table1_rows = [
        ['Date Range', 'Phase / Milestone', 'Activities & Deliverables', 'Status', 'Guide Signature'],
        [
            '06/09/2026 – 12/09/2026',
            'Topic Finalisation & Literature Study',
            'Finalised project topic on deepfake detection. Studied 8 research papers (Rossler, Frank, Qian, Afchar, Li, Tan, Selvaraju, Durall). Identified FaceForensics++ and Celeb-DF v2 as benchmark datasets. Set up Python environment: PyTorch 2.0+, timm, OpenCV, SciPy, Streamlit.',
            'Completed',
            ''
        ],
        [
            '13/09/2026 – 19/09/2026',
            'Synopsis Preparation',
            'Prepared formal MINI PROJECT SYNOPSIS in LIET-prescribed format. Defined 6 project objectives. Drafted Literature Survey, Work Plan, Hardware/Software requirements. Submitted draft to Dr. Devendra Gautam for review.',
            'Completed',
            ''
        ],
        [
            '20/09/2026 – 26/09/2026',
            'Synopsis Approval & Architecture Design',
            'Incorporated mentor feedback. Received synopsis approval. Designed dual-stream pipeline diagram, DFD Level-0 and Level-1. Initialized GitHub repository: github.com/DevRanaKaila/DeepGuard-AI. Initial commit with project structure.',
            'Completed',
            ''
        ],
        [
            '27/09/2026 – 03/10/2026',
            'Dataset Curation & Preprocessing Pipeline',
            'Set up FaceForensics++ dataset structure across raw, c23, c40 tiers. Implemented video frame extraction at 5 FPS using OpenCV. Integrated MTCNN for face detection with 15% boundary margin. Verified preprocessing pipeline on 50 sample videos. Created src/data/ module.',
            'Completed',
            ''
        ],
        [
            '04/10/2026 – 10/10/2026',
            'Spatial Stream Implementation',
            'Integrated EfficientNet-B4 backbone via timm (pretrained ImageNet). Implemented SpatialStream module in src/models/spatial_stream.py. Validated 1792-d feature vector output and gradient flow. Spatial-only baseline achieves 82% accuracy.',
            'Completed',
            ''
        ],
        [
            '11/10/2026 – 17/10/2026',
            'DCT Transform & Mid-Term Review',
            'Implemented 2D-DCT-II transform and Radial Azimuthal Spectrum A(r) in src/data/dct_transform.py. Demonstrated spatial pipeline to mentor. Prepared mid-term progress report. Frequency stream CNN prototype initiated. Mid-term review signed off.',
            'Completed',
            ''
        ],
        [
            '18/10/2026 – 24/10/2026',
            'Frequency Stream & Fusion Module',
            'Implemented 4-stage Spectral-CNN (src/models/frequency_stream.py). Implemented attention-gated bilinear fusion (src/models/fusion.py). Complete DualStreamNet integration. Configured training engine with AdamW, label smoothing, and AMP.',
            'Completed',
            ''
        ],
        [
            '25/10/2026 – 31/10/2026',
            'Dashboard, Grad-CAM & Testing',
            'Develop Streamlit web dashboard (app.py). Implement Grad-CAM heatmap generator (src/utils/gradcam.py). End-to-end pipeline integration. Cross-dataset testing on Celeb-DF v2. Benchmark against XceptionNet and F3-Net baselines.',
            'Planned',
            ''
        ],
        [
            '01/11/2026 – 07/11/2026',
            'Report Writing & Code Finalization',
            'Draft complete 60+ page project report strictly adhering to AKTU formatting rules. Conduct ablation study (spatial-only vs. dual-stream). Run plagiarism check (Turnitin < 20%). Finalize GitHub repo with full documentation and README.',
            'Planned',
            ''
        ],
        [
            '08/11/2026 – 14/11/2026',
            'Final Submission & Viva Defense',
            'Spiral-bind project report. Prepare 15-slide viva presentation deck. Live demo of Streamlit dashboard with real and fake test videos. Final submission to Dr. Devendra Gautam. Panel question-and-answer defense.',
            'Planned',
            ''
        ]
    ]
    
    t1 = doc.add_table(len(table1_rows), 5)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row_vals in enumerate(table1_rows):
        is_hdr = (r_idx == 0)
        row = t1.rows[r_idx]
        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            cell.width = t1_widths[c_idx]
            cell.text = val
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(3)
            cp.paragraph_format.space_after = Pt(3)
            cp.paragraph_format.line_spacing = 1.05
            if is_hdr or c_idx in [0, 3, 4]:
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            for run in cp.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(8.5 if not is_hdr else 9.5)
                run.font.bold = is_hdr or (c_idx == 1)
                if c_idx == 3 and val == 'Completed':
                    run.font.color.rgb = RGBColor(0, 128, 0)
                elif c_idx == 3 and val == 'Planned':
                    run.font.color.rgb = RGBColor(180, 100, 0)
                    
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="80" w:type="dxa"/><w:bottom w:w="80" w:type="dxa"/><w:left w:w="100" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tcMar>')
            tcPr.append(tcMar)
            borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="888888\"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="888888\"/><w:left w:val="single" w:sz="4" w:space="0" w:color="888888\"/><w:right w:val="single" w:sz="4" w:space="0" w:color="888888\"/></w:tcBorders>')
            tcPr.append(borders)
            if is_hdr:
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E2E8F0"/>')
                tcPr.append(shd)
                
    # Footer note
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(6)
    p_note.paragraph_format.space_after = Pt(8)
    r_note = p_note.add_run('Note: Weeks marked Completed have been verified and signed off by the Project Guide. Remaining weeks will be evaluated upon milestone completion.')
    r_note.font.name = 'Times New Roman'
    r_note.font.size = Pt(8.5)
    r_note.font.italic = True
    
    # Signature Table
    t2 = doc.add_table(2, 2)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2.rows[0].cells[0].width = Inches(5.15)
    t2.rows[0].cells[1].width = Inches(5.15)
    t2.rows[1].cells[0].width = Inches(5.15)
    t2.rows[1].cells[1].width = Inches(5.15)
    
    p_sig1 = t2.rows[0].cells[0].paragraphs[0]
    p_sig1.paragraph_format.space_before = Pt(12)
    r_s1 = p_sig1.add_run('Student Signature: _______________________')
    r_s1.font.name = 'Times New Roman'
    r_s1.font.size = Pt(10)
    
    p_sig2 = t2.rows[0].cells[1].paragraphs[0]
    p_sig2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig2.paragraph_format.space_before = Pt(12)
    r_s2 = p_sig2.add_run('Guide Signature: _______________________')
    r_s2.font.name = 'Times New Roman'
    r_s2.font.size = Pt(10)
    
    p_nam1 = t2.rows[1].cells[0].paragraphs[0]
    p_nam1.paragraph_format.space_before = Pt(2)
    r_n1 = p_nam1.add_run('Dev Rana Kaila  |  Roll No. 2301531530023')
    r_n1.font.name = 'Times New Roman'
    r_n1.font.size = Pt(10)
    r_n1.font.bold = True
    
    p_nam2 = t2.rows[1].cells[1].paragraphs[0]
    p_nam2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_nam2.paragraph_format.space_before = Pt(2)
    r_n2 = p_nam2.add_run('Dr. Devendra Gautam  |  Assoc. Professor, CSE')
    r_n2.font.name = 'Times New Roman'
    r_n2.font.size = Pt(10)
    r_n2.font.bold = True
    
    for row in t2.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:bottom w:val="none"/><w:left w:val="none"/><w:right w:val="none"/></w:tcBorders>')
            tcPr.append(borders)
            
    out_paths = [
        r'D:\mini-project\SUBMISSION\Weekly_Progress_Log_Dev_Rana_Kaila.docx',
        r'D:\mini-project\Weekly_Progress_Log_Dev_Rana_Kaila.docx'
    ]
    for p in out_paths:
        doc.save(p)
        print(f'Saved DOCX to {p}')
        
    # Convert to PDF
    import win32com.client
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        for p in out_paths:
            pdf_p = p.replace('.docx', '.pdf')
            wb = word.Documents.Open(os.path.abspath(p), ReadOnly=True)
            wb.SaveAs(os.path.abspath(pdf_p), FileFormat=17)
            wb.Close()
            print(f'Saved PDF to {pdf_p}')
    finally:
        word.Quit()

if __name__ == '__main__':
    generate_weekly_log()
